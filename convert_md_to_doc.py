#!/usr/bin/env python3
"""
Convert Markdown file to Word document with proper styling
"""

import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def convert_markdown_to_docx(md_file_path, docx_file_path):
    """Convert markdown file to Word document with styles"""
    
    # Read markdown file
    with open(md_file_path, 'r', encoding='utf-8') as f:
        content = f.read()
    
    # Create Word document
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    
    lines = content.split('\n')
    i = 0
    
    while i < len(lines):
        line = lines[i]
        original_line = line
        line_stripped = line.strip()
        
        # Skip empty lines
        if not line_stripped:
            i += 1
            continue
        
        # Title (single #)
        if line_stripped.startswith('# ') and not line_stripped.startswith('##'):
            title_text = line_stripped[2:].strip()
            title = doc.add_heading(title_text, level=0)
            title.alignment = WD_ALIGN_PARAGRAPH.LEFT
            i += 1
            continue
        
        # Heading 1 (##)
        if line_stripped.startswith('## ') and not line_stripped.startswith('###'):
            heading_text = line_stripped[3:].strip()
            doc.add_heading(heading_text, level=1)
            i += 1
            continue
        
        # Heading 2 (###)
        if line_stripped.startswith('### '):
            heading_text = line_stripped[4:].strip()
            doc.add_heading(heading_text, level=2)
            i += 1
            continue
        
        # Numbered list (handle indentation)
        if re.match(r'^\s*\d+\.', line):
            # Collect all numbered list items at same or deeper indentation
            list_items = []
            base_indent = len(line) - len(line.lstrip())
            
            while i < len(lines):
                current_line = lines[i]
                current_stripped = current_line.strip()
                
                if not current_stripped:
                    break
                
                # Check if it's a numbered list item
                if re.match(r'^\s*\d+\.', current_line):
                    current_indent = len(current_line) - len(current_line.lstrip())
                    if current_indent >= base_indent:
                        item_text = re.sub(r'^\s*\d+\.\s*', '', current_stripped)
                        list_items.append((item_text, current_indent - base_indent))
                        i += 1
                    else:
                        break
                # Check if it's a sub-item (indented bullet)
                elif current_line.strip().startswith('- ') and len(current_line) - len(current_line.lstrip()) > base_indent:
                    item_text = current_line.strip()[2:].strip()
                    list_items.append((item_text, len(current_line) - len(current_line.lstrip()) - base_indent))
                    i += 1
                else:
                    break
            
            # Add numbered list items
            for item_text, indent_level in list_items:
                if indent_level > 0:
                    # Sub-item - use bullet list
                    p = doc.add_paragraph(style='List Bullet 2')
                else:
                    p = doc.add_paragraph(style='List Number')
                add_formatted_text(p, item_text)
            continue
        
        # Bullet list (handle indentation)
        if line_stripped.startswith('- '):
            # Collect all bullet list items
            list_items = []
            base_indent = len(line) - len(line.lstrip())
            
            while i < len(lines):
                current_line = lines[i]
                current_stripped = current_line.strip()
                
                if not current_stripped:
                    break
                
                if current_line.strip().startswith('- '):
                    current_indent = len(current_line) - len(current_line.lstrip())
                    if current_indent >= base_indent:
                        item_text = current_line.strip()[2:].strip()
                        list_items.append((item_text, current_indent - base_indent))
                        i += 1
                    else:
                        break
                else:
                    break
            
            # Add bullet list items
            for item_text, indent_level in list_items:
                if indent_level > 0:
                    p = doc.add_paragraph(style='List Bullet 2')
                else:
                    p = doc.add_paragraph(style='List Bullet')
                add_formatted_text(p, item_text)
            continue
        
        # Regular paragraph
        if line_stripped:
            process_text_with_formatting(doc, line_stripped, is_list=False)
        
        i += 1
    
    # Save document
    doc.save(docx_file_path)
    print(f"Successfully converted {md_file_path} to {docx_file_path}")

def add_formatted_text(paragraph, text):
    """Add text to paragraph with bold formatting"""
    # Handle bold text (**text**)
    parts = re.split(r'(\*\*[^*]+\*\*)', text)
    
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            # Bold text
            bold_text = part[2:-2]
            run = paragraph.add_run(bold_text)
            run.bold = True
        else:
            # Regular text
            if part:
                paragraph.add_run(part)

def process_text_with_formatting(doc, text, is_list=False, list_type='bullet'):
    """Process text with bold formatting and add to document"""
    
    if is_list:
        if list_type == 'bullet':
            p = doc.add_paragraph(style='List Bullet')
        else:
            p = doc.add_paragraph(style='List Number')
    else:
        p = doc.add_paragraph()
    
    add_formatted_text(p, text)

if __name__ == '__main__':
    md_file = 'Remote switching - User Manual.md'
    docx_file = 'Remote switching - User Manual.docx'
    
    try:
        convert_markdown_to_docx(md_file, docx_file)
    except Exception as e:
        print(f"Error: {e}")
        import traceback
        traceback.print_exc()

